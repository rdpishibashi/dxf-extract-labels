"""機器符号一覧のパイプライン組み込み（「機器符号候補」列の DEF/CAN/空欄 3値化）
のエンドツーエンド仕様テスト（v3.10.0）。

**守る対象**: 機器符号一覧（`designator_master.DesignatorMasterIndex`）に
一致する符号は、どの出力経路を通っても最終的に Excel 上で `'DEF'` になること。
単体テストだけでは「`master_index` の渡し漏れでその経路だけ黙ってDEFが消える」
という不具合を検出できない（`Tools/CLAUDE.md` に記録された
`extract_title_option`/`extract_drawing_numbers_option` の実際の事故と同じ
クラス）ため、`model/ref_designator.py` の builder 関数から
`model/excel_output.py`/`model/assembly_compare.py` の出力関数まで、
実際に Excel を生成して openpyxl で読み戻し確認する。

カバーする経路（引き継ぎ書 2026-10-08、表2「呼び出し元×Excel出力関数」）:
  - 通常モード: build_ref_designator_final() → create_ref_designator_excel_output()
  - 領域モード: build_ref_designator_region_results() → create_region_excel_output()
  - 結線図-組立図比較: compare_assembly()（初回抽出・「判定を実行」再生成の
    いずれからも同じ引数で呼ばれるため、経路としては1つに集約される）

`master_index=None`（鍵未設定・ダイジェストファイル無し・鍵不一致のいずれか）
のときは、3値化前と完全に同じ出力（CAN/空欄のみ、DEFは一切出ない）になることも
あわせて固定する。
"""
import os
import sys
import tempfile
from io import BytesIO

import openpyxl
import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.path.insert(0, PROJECT_ROOT)

from model import assembly_compare  # noqa: E402
from model import ref_designator  # noqa: E402
from model.designator_master import (  # noqa: E402
    load_index, MasterResult, MasterRow, write_digest_file,
)
from model.excel_output import (  # noqa: E402
    create_ref_designator_excel_output, create_region_excel_output,
)
from model.ref_designator import MARK_CANDIDATE, MARK_DEFINED, designator_mark  # noqa: E402


def _load_wb(data: bytes):
    return openpyxl.load_workbook(BytesIO(data))


@pytest.fixture
def master_index():
    """CB001・ELB001（元は `ELB(CB)001`）・F023 を登録した合成索引。
    R10/CN3/GND は意図的に含めない（候補パターンのみ一致＝CANになることの
    対照群として使う）。"""
    result = MasterResult(rows=[
        MasterRow('TEST', 'CB001', '用途1', 1),
        MasterRow('TEST', 'ELB(CB)001', '用途2', 2),
        MasterRow('TEST', 'F023', '', 3),
    ])
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, 'digests.txt')
        write_digest_file(result, path, 'test-secret')
        idx = load_index(path, 'test-secret')
        yield idx


# ---------------------------------------------------------------------------
# designator_mark() の3値（表4: ラベル種別 × 一覧の有無）
# ---------------------------------------------------------------------------

@pytest.mark.parametrize('label, expected_with_index, expected_without_index', [
    # 一覧にある（別称(CB)除去・空白混入の実データ表記も含む）
    ('CB001', MARK_DEFINED, MARK_CANDIDATE),
    ('ELB001', MARK_DEFINED, MARK_CANDIDATE),       # ELB(CB)001 → (CB)除去後と一致
    ('ELB001 (5A/30mA)', MARK_DEFINED, MARK_CANDIDATE),  # 括弧付き実データ表記
    ('F023', MARK_DEFINED, MARK_CANDIDATE),
    # 候補パターンのみ一致（一覧には無い）
    ('R10', MARK_CANDIDATE, MARK_CANDIDATE),
    ('CN3', MARK_CANDIDATE, MARK_CANDIDATE),
    # どちらにも該当しない
    ('(2/5)', None, None),
    ('タイトル', None, None),
])
def test_designator_mark_three_values(master_index, label, expected_with_index, expected_without_index):
    """一覧にある符号は索引ありで'DEF'、索引なしでは'CAN'に下がる
    （候補パターンには一致するため）。候補パターンのみ一致の符号は索引の
    有無に関わらず'CAN'。どちらにも該当しない文字列は索引の有無に関わらずNone。"""
    assert designator_mark(label, master_index) == expected_with_index
    assert designator_mark(label, None) == expected_without_index


# ---------------------------------------------------------------------------
# 通常モード: build_ref_designator_final() → create_ref_designator_excel_output()
# ---------------------------------------------------------------------------

def _normal_mode_ref_data():
    return {
        'a.dxf': {
            'labels': [('CB001', 0, 0), ('R10', 1, 1), ('(2/5)', 2, 2)],
            'total_in_frame': 3,
            'frames': 1,
            'main_drawing_number': 'EE1234-500-01A',
            'title': 'TITLE-A',
            'subtitle': None,
        },
    }


def test_normal_mode_def_reaches_file_sheet_and_total_sheet(master_index):
    """通常モード: master_index渡し時、各ファイルシート・抽出ラベル一覧シートの
    両方にDEFが到達する（渡し漏れがあればどちらかだけDEFが欠ける）。"""
    ref_final = ref_designator.build_ref_designator_final(
        _normal_mode_ref_data(), sort_value='asc', master_index=master_index)
    # Summaryの「機器符号候補ラベル数」= DEF+CANの個数合計（R10のCANも含め2）
    assert ref_final['a.dxf']['ref_designator_count'] == 2

    xlsx = create_ref_designator_excel_output(
        ref_final, 'asc', master_index=master_index)
    wb = _load_wb(xlsx)

    file_ws = wb['a']
    file_rows = {r[1].value: r[0].value for r in file_ws.iter_rows(min_row=2)}
    assert file_rows['CB001'] == 'DEF'
    assert file_rows['R10'] == 'CAN'
    assert file_rows['(2/5)'] is None

    total_ws = wb['抽出ラベル一覧']
    total_rows = {r[1].value: r[0].value for r in total_ws.iter_rows(min_row=2)}
    assert total_rows['CB001'] == 'DEF'
    assert total_rows['R10'] == 'CAN'

    summary_ws = wb['Summary']
    header = [c.value for c in summary_ws[1]]
    count_col = header.index('機器符号候補ラベル数')
    assert summary_ws.cell(row=2, column=count_col + 1).value == 2


def test_normal_mode_without_master_index_matches_legacy_output():
    """master_index=None（鍵未設定時の既定動作）では、DEFが一切出ず、
    3値化前と同じ CAN/空欄のみの出力になる。"""
    ref_final = ref_designator.build_ref_designator_final(
        _normal_mode_ref_data(), sort_value='asc', master_index=None)
    xlsx = create_ref_designator_excel_output(ref_final, 'asc', master_index=None)
    wb = _load_wb(xlsx)

    all_marks = set()
    for ws_name in ('a', '抽出ラベル一覧'):
        for row in wb[ws_name].iter_rows(min_row=2):
            all_marks.add(row[0].value)
    assert MARK_DEFINED not in all_marks
    assert all_marks <= {MARK_CANDIDATE, None}


# ---------------------------------------------------------------------------
# 領域モード: build_ref_designator_region_results() → create_region_excel_output()
# ---------------------------------------------------------------------------

def _region_mode_fixtures():
    square = [(0.0, 0.0), (100.0, 0.0), (100.0, 100.0), (0.0, 100.0)]
    ref_data_by_file = {
        'a.dxf': {
            'labels': [('CB001', 10, 10), ('R10', 20, 20)],
            'total_in_frame': 2,
        },
    }
    region_analyses = {
        'a.dxf': {
            'frames': [(0, 0, 0, 0)],
            'regions': [{'id': 0, 'frame': 0, 'polygon': square, 'area_pct': 50.0,
                         'name_candidates': [(1.0, 'RACK1')]}],
            'main_drawing_number': 'EE1234-500-01A',
            'title': 'TITLE-A',
            'subtitle': None,
        },
    }
    name_selections_by_file = {'a.dxf': {('a.dxf', 0): ['RACK1']}}
    return ref_data_by_file, region_analyses, name_selections_by_file


def test_region_mode_def_reaches_all_four_sheets(master_index):
    """領域モード: master_index渡し時、各ファイルシート・抽出ラベル一覧・
    領域一覧（領域内機器符号候補数）・領域別ラベル一覧の4か所すべてにDEFが
    （または合算件数として）到達する。"""
    ref_data_by_file, region_analyses, name_selections_by_file = _region_mode_fixtures()
    region_results = ref_designator.build_ref_designator_region_results(
        ref_data_by_file, region_analyses, name_selections_by_file, sort_value='asc',
        master_index=master_index)
    # ref_designator_count = DEF(CB001)+CAN(R10) の個数合計 = 2
    assert region_results['a.dxf']['ref_designator_count'] == 2
    # 領域一覧用: region_rows の ref_designator_count もDEF+CAN合算で2
    assert region_results['a.dxf']['region_rows'][0]['ref_designator_count'] == 2

    xlsx = create_region_excel_output(region_results, master_index=master_index)
    wb = _load_wb(xlsx)

    file_ws = wb['a']
    file_rows = {r[1].value: r[0].value for r in file_ws.iter_rows(min_row=2)}
    assert file_rows['CB001'] == 'DEF'
    assert file_rows['R10'] == 'CAN'

    total_ws = wb['抽出ラベル一覧']
    total_rows = {r[1].value: r[0].value for r in total_ws.iter_rows(min_row=2)}
    assert total_rows['CB001'] == 'DEF'

    region_ws = wb['領域一覧']
    header = [c.value for c in region_ws[1]]
    count_col = header.index('領域内機器符号候補数')
    # DEF(CB001)+CAN(R10)の合算で2（3値化してもこの列の値は不変）
    assert region_ws.cell(row=2, column=count_col + 1).value == 2

    lbl_ws = wb['領域別ラベル一覧']
    lbl_rows = {r[2].value: r[1].value for r in lbl_ws.iter_rows(min_row=2)}
    assert lbl_rows['CB001'] == 'DEF'
    assert lbl_rows['R10'] == 'CAN'


def test_region_mode_without_master_index_matches_legacy_output():
    """領域モードでも master_index=None なら DEF が一切出ない。"""
    ref_data_by_file, region_analyses, name_selections_by_file = _region_mode_fixtures()
    region_results = ref_designator.build_ref_designator_region_results(
        ref_data_by_file, region_analyses, name_selections_by_file, sort_value='asc',
        master_index=None)
    xlsx = create_region_excel_output(region_results, master_index=None)
    wb = _load_wb(xlsx)

    all_marks = set()
    for ws_name in ('a', '抽出ラベル一覧', '領域別ラベル一覧'):
        for row in wb[ws_name].iter_rows(min_row=2):
            all_marks.add(row[0].value if ws_name != '領域別ラベル一覧' else row[1].value)
    assert MARK_DEFINED not in all_marks


# ---------------------------------------------------------------------------
# 結線図-組立図比較: compare_assembly()
# ---------------------------------------------------------------------------

def _assembly_compare_results():
    return {
        'wiring.dxf': {
            'rows': [{'機器符号候補': None, 'ラベル': 'CB001', '個数': 2},
                     {'機器符号候補': None, 'ラベル': 'R10', '個数': 1}],
            'title': 'UNIT内結線図', 'main_drawing_number': 'EE001',
        },
        'assembly.dxf': {
            'rows': [{'機器符号候補': None, 'ラベル': 'CB001', '個数': 2}],
            'title': '組立図', 'main_drawing_number': 'EE002',
        },
    }


def test_compare_assembly_def_reaches_comparison_sheet(master_index):
    """結線図-組立図比較: master_index渡し時、比較シートの行にもDEFが到達する
    （初回抽出〈_resolve_assembly_compare〉・「判定を実行」再生成
    〈_apply_type_judgment〉のいずれから呼ばれても同一引数のため経路は1つ）。"""
    results = _assembly_compare_results()
    compare_result = assembly_compare.compare_assembly(
        results, ['wiring.dxf'], 'assembly.dxf', master_index=master_index)
    rows_by_label = {r['ラベル']: r for r in compare_result['rows']}
    assert rows_by_label['CB001']['機器符号候補'] == 'DEF'
    assert rows_by_label['R10']['機器符号候補'] == 'CAN'

    ref_final = ref_designator.build_ref_designator_final(
        _normal_mode_ref_data(), sort_value='asc', master_index=master_index)
    xlsx = create_ref_designator_excel_output(
        ref_final, 'asc', assembly_compare=compare_result, master_index=master_index)
    wb = _load_wb(xlsx)
    cmp_ws = wb['結線図-組立図比較']
    cmp_rows = {r[1].value: r[0].value for r in cmp_ws.iter_rows(min_row=2)}
    assert cmp_rows['CB001'] == 'DEF'
    assert cmp_rows['R10'] == 'CAN'


def test_compare_assembly_without_master_index_matches_legacy_output():
    """compare_assembly()もmaster_index=Noneなら従来通りCAN/空欄のみ。"""
    results = _assembly_compare_results()
    compare_result = assembly_compare.compare_assembly(
        results, ['wiring.dxf'], 'assembly.dxf', master_index=None)
    marks = {r['機器符号候補'] for r in compare_result['rows']}
    assert MARK_DEFINED not in marks
